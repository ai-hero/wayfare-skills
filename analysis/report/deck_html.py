"""Read a findings deck (.pptx) back into what it says, for the book's pages.

The deck is the source: native chart data, the side text, milestone lines and speaker notes are
read back out of the .pptx, so no deck.py changes to get a page. Slides that share a tag prefix
("Q work-sources · Answer", "Q work-sources · By repo", "Q work-sources · Change sets") become
tabs of one question, so a deck adds a view by adding a slide with the same prefix.

book_pages.py renders the book with this; `fill` injects a page's content, as JSON, into the
prebuilt React viewer (shadcn charts on Recharts), whose build is the only step that needs node.
"""
import base64
import html
import json
import os
import re
import sys
from datetime import date, timedelta

from pptx import Presentation

from pptx.enum.shapes import MSO_SHAPE_TYPE

EMU = 914400
C = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
# The viewer is built from the private design system's components, so it lives in the gitignored
# .analysis/ beside the data rather than in this public repo.
VIEWER = os.environ.get("FINDINGS_VIEWER", os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(
    os.path.abspath(__file__)))), ".analysis", "viewer", "dist", "index.html"))
SIDE_X = 9.0
MONTHS = {m: i for i, m in enumerate(["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov",
                                      "Dec"], 1)}


def inch(v):
    return (v or 0) / EMU


def text_of(el):
    return "".join(t.text or "" for t in el.iter(A + "t"))


def paragraphs(shape):
    return [p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip()]


# ------------------------------------------------------------------ charts

def _pts(cache, n=None):
    if cache is None:
        return []
    count = cache.find(C + "ptCount")
    n = int(count.get("val")) if count is not None else n or 0
    out = [None] * n
    for pt in cache.findall(C + "pt"):
        i = int(pt.get("idx"))
        v = pt.find(C + "v")
        if i >= len(out):
            out.extend([None] * (i + 1 - len(out)))
        out[i] = v.text if v is not None else None
    return out


def _color(sppr):
    if sppr is None:
        return None
    for tag in ("solidFill", "ln"):
        el = sppr.find(A + tag)
        if el is None:
            continue
        clr = el.find(".//" + A + "srgbClr")
        if clr is not None:
            return clr.get("val")
    return None


def _axis_title(ax):
    t = ax.find(C + "title") if ax is not None else None
    return text_of(t).strip() if t is not None else None


def series_name(tx):
    if tx is None:
        return ""
    v = tx.find(f".//{C}v")
    return (v.text if v is not None else text_of(tx)).strip()


def _weekly_labels(cats):
    """Long timelines label every 4th week and blank the rest; put the week back on every one."""
    known = [(i, c) for i, c in enumerate(cats) if re.fullmatch(r"\d{1,2} [A-Z][a-z]{2}", c or "")]
    if not known or all((c or "").strip() for c in cats):
        return cats
    i0, c0 = known[0]
    d, m = c0.split()
    y = 2025 if MONTHS[m] == 12 and i0 == 0 else 2026
    base = date(y, MONTHS[m], int(d)) - timedelta(days=7 * i0)
    return [(base + timedelta(days=7 * i)).strftime("%-d %b") for i in range(len(cats))]


def read_chart(shape, where=""):
    cs = shape.chart._chartSpace
    plot_area = cs.find(f"{C}chart/{C}plotArea")
    ml = plot_area.find(f"{C}layout/{C}manualLayout")
    inner = {k: float(ml.find(C + k).get("val")) for k in ("x", "y", "w", "h")} if ml is not None else \
        {"x": 0.09, "y": 0.16, "w": 0.88, "h": 0.68}
    plot = next((el for el in plot_area if el.tag in (C + "barChart", C + "lineChart")), None)
    if plot is None:
        raise ValueError(f"{where}: only bar and line charts can be rendered, not {plot_area[-1].tag.split('}')[-1]}")
    kind = "line" if plot.tag == C + "lineChart" else "bar"
    horizontal = kind == "bar" and plot.find(C + "barDir").get("val") == "bar"
    grouping = plot.find(C + "grouping")
    stacked = grouping is not None and grouping.get("val") in ("stacked", "percentStacked")
    series, cats = [], None
    for ser in plot.findall(C + "ser"):
        cat = ser.find(C + "cat")
        if cats is None and cat is not None:
            cache = cat.find(f".//{C}strCache")
            if cache is None:
                cache = cat.find(f".//{C}numCache")
            cats = [c or "" for c in _pts(cache)]
        vals = _pts(ser.find(f"{C}val//{C}numCache"), len(cats or []))
        points = {int(dp.find(C + "idx").get("val")): _color(dp.find(C + "spPr")) for dp in ser.findall(C + "dPt")}
        series.append({
            "name": series_name(ser.find(C + "tx")),
            "values": [None if v is None else float(v) for v in vals],
            "color": _color(ser.find(C + "spPr")),
            "pointColors": [points.get(i) for i in range(len(vals))] if points else None,
        })
    dl = plot.find(C + "dLbls")
    labels = dl is not None and dl.find(C + "showVal") is not None and dl.find(C + "showVal").get("val") == "1"
    val_ax, cat_ax = plot_area.find(C + "valAx"), plot_area.find(C + "catAx")
    fmt = val_ax.find(C + "numFmt").get("formatCode") if val_ax.find(C + "numFmt") is not None else "General"
    scaling = val_ax.find(C + "scaling")
    vmax = scaling.find(C + "max")
    vmin = scaling.find(C + "min")
    cats = _weekly_labels(cats or [])
    if horizontal:
        cats = cats[::-1]
        for s in series:
            s["values"] = s["values"][::-1]
            if s["pointColors"]:
                s["pointColors"] = s["pointColors"][::-1]
    return {
        "kind": kind, "horizontal": horizontal, "stacked": stacked, "cats": cats, "series": series,
        "labels": labels, "fmt": fmt, "legend": cs.find(f"{C}chart/{C}legend") is not None,
        "xTitle": _axis_title(cat_ax), "yTitle": _axis_title(val_ax),
        "max": float(vmax.get("val")) if vmax is not None else None,
        "min": float(vmin.get("val")) if vmin is not None else None,
        "frame": (inch(shape.left), inch(shape.top), inch(shape.width), inch(shape.height)), "inner": inner,
        "marks": [], "bands": [],
    }


def attach_overlays(charts, lines, rotated, rects):
    """Milestone lines and shaded spans are drawn shapes; turn each into a category position on its chart."""
    def owner(x):
        for ch in charts:
            fx, _, fw, _ = ch["frame"]
            if fx <= x <= fx + fw and not ch["horizontal"]:
                return ch
        return None

    def units(ch, x):
        fx, _, fw, _ = ch["frame"]
        left = fx + fw * ch["inner"]["x"]
        return (x - left) / (fw * ch["inner"]["w"]) * len(ch["cats"])

    used = set()
    for ln in lines:
        ch = owner(ln["x"])
        if not ch:
            continue
        near = min(((abs(t["cx"] - ln["x"]), i) for i, t in enumerate(rotated) if i not in used), default=None)
        name = ""
        if near and near[0] < 0.2:
            used.add(near[1])
            name = rotated[near[1]]["text"]
        ch["marks"].append({"at": round(units(ch, ln["x"]), 3), "name": name, "color": ln["color"]})
    for r in rects:
        ch = owner(r["x"] + r["w"] / 2)
        if ch:
            ch["bands"].append({"from": round(units(ch, r["x"]), 3), "to": round(units(ch, r["x"] + r["w"]), 3),
                                "name": r["text"], "color": r["color"]})


# ------------------------------------------------------------------ slides

def read_table(shape):
    return [[cell.text.strip() for cell in row.cells] for row in shape.table.rows]


def read_slide(slide, where=""):
    s = {"layout": slide.slide_layout.name, "charts": [], "tables": [], "images": [], "texts": [], "notes": ""}
    lines, rotated, rects = [], [], []
    for sh in slide.shapes:
        x, y, w, h = inch(sh.left), inch(sh.top), inch(sh.width), inch(sh.height)
        if sh.has_chart:
            s["charts"].append(read_chart(sh, where))
        elif sh.shape_type == MSO_SHAPE_TYPE.TABLE:
            s["tables"].append(read_table(sh))
        elif sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
            img = sh.image
            s["images"].append({"src": f"data:{img.content_type};base64,{base64.b64encode(img.blob).decode()}",
                                "alt": sh.name})
        elif sh.shape_type == MSO_SHAPE_TYPE.LINE:
            try:
                col = str(sh.line.color.rgb)
            except (AttributeError, TypeError):
                col = None
            lines.append({"x": x, "color": col})
        elif sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
            try:
                col = str(sh.fill.fore_color.rgb)
            except (AttributeError, TypeError):
                col = None
            rects.append({"x": x, "w": w, "text": sh.text_frame.text.strip() if sh.has_text_frame else "",
                          "color": col})
        elif sh.has_text_frame and sh.text_frame.text.strip():
            if abs((sh.rotation or 0) - 270) < 1:
                rotated.append({"cx": x + w / 2, "text": sh.text_frame.text.strip()})
            else:
                s["texts"].append({"x": x, "y": y, "w": w, "paras": paragraphs(sh),
                                   "placeholder": sh.is_placeholder})
    attach_overlays(s["charts"], lines, rotated, rects)
    s["texts"].sort(key=lambda t: (t["y"], t["x"]))
    if slide.has_notes_slide:
        s["notes"] = slide.notes_slide.notes_text_frame.text.strip()
    return s


def classify(s):
    """Turn a slide's shapes into what it says: its kind, its tag and its parts."""
    t = s["texts"]
    if s["layout"] == "Title":
        return {"kind": "title", "title": t[0]["paras"][0] if t else "",
                "subtitle": " ".join(t[1]["paras"]) if len(t) > 1 else "", "notes": s["notes"]}
    if s["layout"] == "Section divider":
        return {"kind": "section", "kicker": " ".join(t[0]["paras"]) if t else "",
                "title": " ".join(t[1]["paras"]) if len(t) > 1 else "", "notes": s["notes"]}
    if s["charts"] or s["tables"] or s["images"]:
        side = [x for x in t if x["x"] >= SIDE_X]
        rest = [x for x in t if x["x"] < SIDE_X]
        tag = side[0]["paras"][0] if side else ""
        source = next((" ".join(x["paras"]) for x in side if x["paras"][0].upper().startswith("SOURCE")), "")
        body = [x for x in side[1:] if not x["paras"][0].upper().startswith("SOURCE")]
        return {"kind": "view", "tag": tag, "title": " ".join(body[0]["paras"]) if body else "",
                "points": [p for x in body[1:] for p in x["paras"]],
                "source": re.sub(r"^SOURCE\s*", "", source, flags=re.I),
                "extra": [p for x in rest for p in x["paras"]],
                "charts": s["charts"], "tables": s["tables"], "images": s["images"], "notes": s["notes"]}
    tag = t[0]["paras"][0] if t else ""
    if tag.upper().endswith("· QUESTION"):
        how, originally, cur = [], [], None
        for x in t[2:]:
            for p in x["paras"]:
                if p.upper() == "HOW WE ANSWER IT":
                    cur = how
                elif p.upper() == "ORIGINALLY ASKED":
                    cur = originally
                elif cur is not None:
                    cur.append(p)
        return {"kind": "question", "tag": tag, "question": " ".join(t[1]["paras"]) if len(t) > 1 else "",
                "how": " ".join(how), "originally": " ".join(originally), "notes": s["notes"]}
    return {"kind": "statement", "tag": tag, "headline": " ".join(t[1]["paras"]) if len(t) > 1 else "",
            "body": [p for x in t[2:] for p in x["paras"]], "notes": s["notes"]}


def split_tag(tag):
    parts = [p.strip() for p in tag.split("·")]
    if len(parts) == 1:
        return tag.strip(), ""
    return parts[0], " · ".join(parts[1:])


def group(slides):
    """Sections holding blocks; a block is one question (or topic) with its views as tabs."""
    head, sections, block = None, [], None
    for s in slides:
        k = s["kind"]
        if k == "title":
            head = s
            continue
        if k == "section" or not sections:
            sections.append({"kicker": s.get("kicker", ""), "title": s.get("title", ""),
                             "notes": s.get("notes", ""), "blocks": []})
            block = None
            if k == "section":
                continue
        blocks = sections[-1]["blocks"]
        if k == "question":
            key, _ = split_tag(s["tag"])
            block = {"key": key, "question": s, "views": [], "statements": []}
            blocks.append(block)
        elif k == "view":
            key, label = split_tag(s["tag"])
            if not block or block["key"].upper() != key.upper():
                block = {"key": key, "question": None, "views": [], "statements": []}
                blocks.append(block)
            s["label"] = label or "Answer"
            block["views"].append(s)
        else:
            blocks.append({"key": s["tag"], "question": None, "views": [], "statements": [s]})
            block = None
    return head or {"title": "", "subtitle": ""}, sections



# ------------------------------------------------------------------ lessons

# Deck eyebrows are set in capitals (Q WORK-SOURCES · QUESTION), so an id matches in any case.
Q_RE = re.compile(r"\bQ\s((?i:[a-z][a-z0-9]*(?:-[a-z0-9]+)+))\b")
VERDICT_RE = re.compile(r"^([A-Z][A-Z', ]+?)\s{2}·\s{2}(.*)$", re.S)


def qkey(v):
    m = Q_RE.search(v or "")
    return f"Q {m.group(1).lower()}" if m else None


def tone(verdict):
    v = verdict.upper()
    if v.startswith("HELD"):
        return "held"
    if v.startswith("PARTLY"):
        return "partly"
    if v.startswith("DID NOT"):
        return "not"
    return "open"


def read_lessons(path):
    """{"Q work-sources": {"insights": [...], "hypotheses": [...]}} from the evidence-edition talk deck."""
    out = {}
    def at(k):
        return out.setdefault(k, {"insights": [], "hypotheses": []})
    for slide in Presentation(path).slides:
        texts = [sh.text_frame for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
        for i, tf in enumerate(texts):
            t = tf.text.strip()
            if t.startswith("HYPOTHESIS") and i + 1 < len(texts):
                paras = [p.text.strip() for p in texts[i + 1].paragraphs if p.text.strip()]
                m = VERDICT_RE.match(paras[-1]) if len(paras) > 1 else None
                if not m:
                    continue
                theme = t.split("\n", 1)[1].strip().title() if "\n" in t else ""
                evidence = re.split(r"\s+·\s+(?=Ch \d)", m.group(2).split("  ·  ")[0])[0].strip()
                refs = sorted({qkey(r.group(0)) for r in Q_RE.finditer(m.group(2))})
                h = {"statement": " ".join(paras[:-1]), "verdict": m.group(1).strip(), "tone": tone(m.group(1)),
                     "evidence": evidence, "refs": refs, "theme": theme}
                for r in refs:
                    at(r)["hypotheses"].append(h)
            if t.upper().endswith("· FINDING"):
                key = qkey(t) or qkey(notes)
                body = next((x for x in texts if "THE INSIGHT" in x.text), None)
                if not key or body is None:
                    continue
                paras = [p.text.strip() for p in body.paragraphs if p.text.strip()]
                if "THE INSIGHT" not in paras:
                    continue
                insight = " ".join(paras[paras.index("THE INSIGHT") + 1:])
                if insight and insight not in at(key)["insights"]:
                    at(key)["insights"].append(insight)
    return out


def attach_lessons(sections, lessons):
    hits = 0
    for s in sections:
        for b in s["blocks"]:
            key = qkey(b["key"])
            if key and key in lessons:
                b["lessons"] = lessons[key]
                hits += 1
    return hits

# ------------------------------------------------------------------ html

def esc(v):
    return html.escape(v or "", quote=True)


def slim(sections):
    """Drop the slide geometry the page doesn't need once overlays are placed."""
    for s in sections:
        for b in s["blocks"]:
            for v in b["views"]:
                for ch in v["charts"]:
                    ch.pop("frame", None)
                    ch.pop("inner", None)
    return sections


def fill(template, title, payload):
    missing = [m for m in ("__DECK_TITLE__", "__DECK_DATA__") if m not in template]
    if missing:
        sys.exit(f"the viewer build has no {', '.join(missing)} marker: rebuild it with `npm run build`")
    data = json.dumps(payload, separators=(",", ":"), ensure_ascii=False, allow_nan=False).replace("</", "<\\/")
    # str.replace, not format(): the bundled script is full of braces.
    return template.replace("__DECK_TITLE__", esc(title), 1).replace("__DECK_DATA__", data, 1)


# ------------------------------------------------------------------ book

FIG_RE = re.compile(r"^\[\[(.+?)\]\]$")
LIST_RE = re.compile(r"^([-*]|\d+\.)\s+(.+)$")
NUM_RE = re.compile(r"\$?\d[\d,]*(?:\.\d+)?%?")
# Numbers a reader can check without the data: question and chapter numbers, dates, years, URLs, model versions,
# and names that are digits (8090 is a company).
FREE_RE = re.compile(r"https?://\S+|\bQ\s(?i:[a-z][a-z0-9]*(?:-[a-z0-9]+)+)\b|\b(?:Chapters?|Ch|Part|Figures?)\s+\d+(?:\.\d+)?(?!,\d)(?:\s*(?:,|and|to|–|-)\s*\d+(?:\.\d+)?(?!,?\d))*"
                     r"|\b\d{1,2}\s+(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*\b"
                     r"|\b(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*\s+\d{1,2}\b|\b(?:19|20)\d\ds?\b"
                     r"|\b(?:Opus|Sonnet|Haiku|Fable|Mythos)\s+\d(?:\.\d)?\b|\bClaude\s+\d(?:\.\d)?(?=\s+(?:Opus|Sonnet|Haiku))|\b8090(?:\.ai)?\b")


WARNINGS = []


def warn(msg):
    WARNINGS.append(msg)
    print(msg, file=sys.stderr)


def parse_book(text):
    """book.md: `## ` sections, `### ` subheads, `> ` callouts, `| a | b |` tables, `- ` and `1. ` lists,
    `[[Q work-sources · By repo | caption]]` figures. `[[I work-sources | caption]]` is the same card in its
    in-text form (question, finding, chart); `Q` places the whole card, which is what the appendix wants."""
    sections, para = [{"title": "", "blocks": []}], []

    def flush():
        if para:
            sections[-1]["blocks"].append({"kind": "p", "text": " ".join(para)})
            para.clear()

    for line in re.sub(r"<!--.*?-->", "", text, flags=re.S).splitlines():
        line = line.strip()
        fig = FIG_RE.match(line)
        if line.startswith("# "):
            continue
        if line.startswith("## "):
            flush()
            sections.append({"title": line[3:].strip(), "blocks": []})
        elif not line:
            flush()
        elif line.startswith("### "):
            flush()
            sections[-1]["blocks"].append({"kind": "h3", "text": line[4:].strip()})
        elif line.startswith(">"):
            flush()
            quote = line.lstrip("> ").strip()
            blocks = sections[-1]["blocks"]
            if blocks and blocks[-1]["kind"] == "quote" and line != ">":
                blocks[-1]["text"] += " " + quote
            elif quote:
                blocks.append({"kind": "quote", "text": quote})
        elif fig:
            flush()
            ref, _, caption = fig.group(1).partition("|")
            ref, inline = re.subn(r"^I\s+", "Q ", ref.strip())
            sections[-1]["blocks"].append({"kind": "figure", "ref": ref, "caption": caption.strip(),
                                           "inline": bool(inline)})
        elif line.startswith("|"):
            flush()
            cells = [c.strip() for c in line.strip("|").split("|")]
            blocks = sections[-1]["blocks"]
            if blocks and blocks[-1]["kind"] == "table":
                # The `| :--- |` rule under the header row is layout, not a row.
                if not all(re.fullmatch(r":?-+:?", c) for c in cells):
                    blocks[-1]["rows"].append(cells)
            else:
                blocks.append({"kind": "table", "rows": [cells]})
        elif LIST_RE.match(line):
            flush()
            m = LIST_RE.match(line)
            kind = "ol" if m.group(1)[0].isdigit() else "ul"
            blocks = sections[-1]["blocks"]
            if blocks and blocks[-1]["kind"] == kind:
                blocks[-1]["items"].append(m.group(2))
            else:
                blocks.append({"kind": kind, "items": [m.group(2)]})
        else:
            para.append(line)
    flush()
    return sections


def find_view(sections, ref):
    """`Q work-sources` is its first view; `Q work-sources · By repo` the view with that label; a diagram by its tag or label."""
    key, label = split_tag(ref)
    want = ref.upper()
    for s in sections:
        for b in s["blocks"]:
            for v in b["views"]:
                vlabel = (v.get("label") or "").upper()
                if qkey(ref):
                    hit = qkey(b["key"]) == qkey(key) and (not label or vlabel == label.upper())
                else:
                    hit = want in (vlabel, f"{b['key']} · {vlabel}".upper(), b["key"].upper())
                if hit:
                    return b, v
    return None, None


def bare(num):
    return num.strip("$%").replace(",", "")


def rounds_to(num, known):
    """`about 695,000` in the prose against 695,450 on the slide: the prose may round to its trailing zeros."""
    b = bare(num)
    if not b.isdigit() or not b.endswith("0"):
        return False
    unit = 10 ** (len(b) - len(b.rstrip("0")))
    return any(k.isdigit() and round(int(k) / unit) * unit == int(b) for k in known)


def haystack(sections):
    parts = []
    for s in sections:
        for b in s["blocks"]:
            q = b.get("question") or {}
            parts += [q.get("question", ""), q.get("how", ""), q.get("notes", "")]
            for v in b["views"]:
                parts += [v["title"], v["source"], v["notes"], *v["points"], *v["extra"]]
                parts += [" ".join(r) for t in v["tables"] for r in t]
            for st in b["statements"]:
                parts += [st["headline"], st["notes"], *st["body"]]
    text = " ".join(parts)
    nums = {bare(n) for n in NUM_RE.findall(text)}
    # A slide writes 456K where the prose writes 456,000: the same number, so both spellings count.
    for n, unit in re.findall(r"\b(\d+(?:\.\d+)?)([KkMm])\b", text):
        nums.add(str(int(float(n) * (1000 if unit in "Kk" else 1000000))))
    return nums


if __name__ == "__main__":
    sys.exit("deck_html.py is a library now; render the book with report/book_pages.py --out DIR")
