"""Build the per-question findings deck from bank/questions.csv, grounded
in the analysis cube, using the AI Hero Deck Template.

Run with the venv that has python-pptx installed, e.g.:
    /tmp/pptxenv/bin/python3 report/build_deck.py <template.pptx> <out.pptx>

One "Statement" (question), "Two column" (analysis) and "Title and body"
(findings) slide per bank row, chaptered with "Section divider" slides.
Layout indices are read by name, not hardcoded, so a template reorder
doesn't silently misfile content into the wrong layout.
"""
import csv
import importlib
import os
import pkgutil
import sys

from pptx import Presentation
from pptx.util import Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ANALYSIS_ROOT = os.path.dirname(HERE)
sys.path.insert(0, ANALYSIS_ROOT)

from cube.db import connect  # noqa: E402

BANK = os.path.join(ANALYSIS_ROOT, "bank", "questions.csv")
QUESTIONS_DIR = os.path.join(ANALYSIS_ROOT, "questions")

# Parsed once from the source deck's own "Section divider" slides
# (Software Factory Research Questions (review).pptx) -- not re-derivable
# from bank/questions.csv, which only carries the topic.
CHAPTER_INFO = {
    "method": ("FRONT MATTER", "Research setting, method and the baseline"),
    "agent_harness": ("PART I", "The harness"),
    "owner": ("PART I", "Human in the loop"),
    "skills": ("PART I", "Skills and factory evolution"),
    "connectors": ("PART II", "Connectors"),
    "memory": ("PART II", "Knowledge and memory"),
    "architecture": ("PART II", "Architecture and design records"),
    "work_items": ("PART III", "Work items, flow and wall time"),
    "mistakes": ("PART III", "Agent mistakes and rework"),
    "security": ("PART III", "Security"),
    "fleet": ("PART IV", "Fleet scope and apps"),
    "messages": ("PART IV", "Cross-repo context and messaging"),
    "compliance": ("PART IV", "Compliance and drift"),
    "deployment": ("PART IV", "Deployment and infrastructure"),
    "spend": ("PART V", "Spend and cost"),
    "efficiency": ("PART V", "Factory floor efficiency"),
    "manager": ("PART V", "The factory manager's thinking"),
    "field": ("PART VI", "The field, and what generalizes"),
}


def load_bank():
    with open(BANK) as f:
        return list(csv.DictReader(f))


def discover_modules():
    by_rq_id = {}
    for _, name, _ in pkgutil.iter_modules([QUESTIONS_DIR]):
        if name in ("runner", "__init__"):
            continue
        mod = importlib.import_module(f"questions.{name}")
        rq_id = getattr(mod, "RQ_ID", None)
        if rq_id:
            by_rq_id[rq_id] = mod
    return by_rq_id


def layouts_by_name(prs):
    out = {}
    for master in prs.slide_masters:
        for layout in master.slide_layouts:
            out[layout.name] = layout
    return out


def _set_size(shape, pt):
    for para in shape.text_frame.paragraphs:
        for run in para.runs:
            run.font.size = Pt(pt)


def set_text(shape, text, max_chars=None, size_pt=None):
    text = text or ""
    if max_chars and len(text) > max_chars:
        text = text[: max_chars - 1].rstrip() + "…"
    shape.text_frame.text = text
    if size_pt:
        _set_size(shape, size_pt)


def set_bullets(shape, lines, max_lines=8, max_chars_per_line=140, size_pt=None):
    tf = shape.text_frame
    tf.clear()
    lines = [l for l in lines if l][:max_lines]
    if not lines:
        lines = ["—"]
    for i, line in enumerate(lines):
        if len(line) > max_chars_per_line:
            line = line[: max_chars_per_line - 1].rstrip() + "…"
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
    if size_pt:
        _set_size(shape, size_pt)


def ph(slide, idx):
    for shape in slide.placeholders:
        if shape.placeholder_format.idx == idx:
            return shape
    return None


def tag_line(row):
    return f"{row['chapter'].upper()} · {row['rq_id']} · {row['method'].upper()} · {row['status'].upper()}"


def run_answer(mod, con):
    try:
        rows = mod.answer(con)
        return [dict(r) if hasattr(r, "keys") else r for r in rows], None
    except Exception as e:  # pragma: no cover - defensive: deck build must not abort mid-run
        return None, f"{type(e).__name__}: {e}"


def findings_lines(rows, err):
    if err:
        return [f"Runtime error while grounding this slide: {err}"]
    if not rows:
        return ["No rows returned."]
    caveat = None
    data_rows = []
    for r in rows:
        if "answerable_by_code" in r:
            caveat = r
        else:
            data_rows.append(r)
    lines = []
    if caveat is not None:
        verdict = caveat.get("answerable_by_code")
        reason = caveat.get("reason") or caveat.get("needs") or ""
        prefix = "PARTIAL" if verdict == "partial" else ("NO" if verdict is False else str(verdict))
        lines.append(f"Answerable from code: {prefix}. {reason}")
    row_budget = 4 if caveat is None else 3
    for r in data_rows[:row_budget]:
        lines.append("; ".join(f"{k}={v}" for k, v in list(r.items())[:5]))
    if not data_rows and caveat is None:
        for r in rows[:row_budget]:
            lines.append("; ".join(f"{k}={v}" for k, v in list(r.items())[:5]))
    return lines


def add_question_slide(prs, layouts, row):
    slide = prs.slides.add_slide(layouts["Statement"])
    set_text(ph(slide, 10), tag_line(row), max_chars=90)
    # The layout's "Number" placeholder defaults to 100pt for a headline
    # stat like "104K" -- a full question sentence needs a much smaller
    # override or it overruns the box (LibreOffice doesn't recompute the
    # template's normAutofit shrink-to-fit on unattended conversion).
    question = row["question"]
    size_pt = 30 if len(question) <= 110 else 24 if len(question) <= 170 else 20
    set_text(ph(slide, 0), question, max_chars=220, size_pt=size_pt)
    ctx = row.get("background") or "Grounded against the fleet's ingested git/GitHub/plan/knowledge data."
    set_text(ph(slide, 1), ctx, max_chars=200, size_pt=14)
    return slide


def add_analysis_slide(prs, layouts, row):
    slide = prs.slides.add_slide(layouts["Two column"])
    set_text(ph(slide, 10), tag_line(row), max_chars=90)
    set_text(ph(slide, 0), "How it's answered", max_chars=90)
    # "Left body" is a single short supporting line in this layout (1.2in
    # tall) -- three terse facts fit, a full hypothesis paragraph doesn't.
    left = [
        f"Method: {row['method']}",
        f"Sources: {(row['sources'] or '—')[:55]}",
        f"Detectors: {(row['detectors'] or '—')[:55]}",
    ]
    set_bullets(ph(slide, 1), left, max_lines=3, max_chars_per_line=60, size_pt=14)
    # "Right body" is the tall column (3.2in) -- the longer prose goes here.
    right = []
    if row.get("hypothesis"):
        right.append(f"Hypothesis: {row['hypothesis']}")
    if row.get("ideal"):
        right.append(f"In an ideal factory: {row['ideal']}")
    if row.get("defer_reason"):
        right.append(f"Deferred because: {row['defer_reason']}")
    if not right:
        right.append("No stated ideal/defer-reason in the bank for this question.")
    set_bullets(ph(slide, 2), right, max_lines=3, max_chars_per_line=150, size_pt=14)
    return slide


def add_findings_slide(prs, layouts, row, rows, err):
    slide = prs.slides.add_slide(layouts["Title and body"])
    set_text(ph(slide, 10), tag_line(row), max_chars=90)
    set_text(ph(slide, 0), "Findings", max_chars=90)
    set_bullets(ph(slide, 1), findings_lines(rows, err), max_lines=5, max_chars_per_line=100, size_pt=12)
    return slide


def add_section_divider(prs, layouts, chapter, questions_in_chapter):
    part, name = CHAPTER_INFO.get(chapter, ("", chapter))
    slide = prs.slides.add_slide(layouts["Section divider"])
    set_text(ph(slide, 10), f"{name.upper()} · {part} · {questions_in_chapter} QUESTIONS", max_chars=90)
    set_text(ph(slide, 0), name, max_chars=90)
    return slide


def add_title_slide(prs, layouts, n_questions):
    slide = prs.slides.add_slide(layouts["Title"])
    set_text(ph(slide, 0), "Architecting a Software Factory", max_chars=90)
    set_text(ph(slide, 1), "Findings report: every research question, grounded in the fleet's own data", max_chars=140)
    set_text(ph(slide, 11), f"{n_questions} questions · generated from analysis/", max_chars=140)
    return slide


def build(template_path, out_path):
    prs = Presentation(template_path)
    layouts = layouts_by_name(prs)
    bank = load_bank()
    modules = discover_modules()
    con = connect()

    add_title_slide(prs, layouts, len(bank))

    current_chapter = None
    chapter_counts = {}
    for row in bank:
        chapter_counts[row["chapter"]] = chapter_counts.get(row["chapter"], 0) + 1

    missing = []
    errors = []
    for row in bank:
        if row["chapter"] != current_chapter:
            current_chapter = row["chapter"]
            add_section_divider(prs, layouts, current_chapter, chapter_counts[current_chapter])

        add_question_slide(prs, layouts, row)
        add_analysis_slide(prs, layouts, row)

        mod = modules.get(row["rq_id"])
        if mod is None:
            missing.append(row["rq_id"])
            rows, err = None, "no questions/*.py module found for this rq_id"
        else:
            rows, err = run_answer(mod, con)
            if err:
                errors.append((row["rq_id"], err))
        add_findings_slide(prs, layouts, row, rows, err)

    prs.save(out_path)
    print(f"wrote {out_path}: {len(list(prs.slides))} slides")
    print(f"bank rows: {len(bank)}, missing modules: {len(missing)}, runtime errors: {len(errors)}")
    if missing:
        print("missing:", missing)
    if errors:
        print("errors:", errors)


if __name__ == "__main__":
    template_path = sys.argv[1] if len(sys.argv) > 1 else "/tmp/sf_deck/template.pptx"
    out_path = sys.argv[2] if len(sys.argv) > 2 else "/tmp/sf_deck/out.pptx"
    build(template_path, out_path)
